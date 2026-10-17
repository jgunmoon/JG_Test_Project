# 테스트용 문서는 코드로 만든 가상의 Chair notes입니다 (실제 3GPP 합의 내용이 아님).
import docx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts import parse_agreements as pa

URL = "https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_999/Inbox/Chair_notes/cn.docx"


def _number(paragraph, ilvl, num_id=1):
    num_pr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", ilvl), ("w:numId", num_id)):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(val))
        num_pr.append(el)
    paragraph._p.get_or_add_pPr().append(num_pr)


class Doc:
    def __init__(self):
        self.d = docx.Document()

    def heading(self, text, ilvl, style="Heading 1", num_id=1):
        _number(self.d.add_paragraph(text, style=style), ilvl, num_id)

    def para(self, text, bold=False):
        self.d.add_paragraph().add_run(text).bold = bold

    def tdoc(self, number, rest, bold=False):
        p = self.d.add_paragraph()
        p.add_run(number).bold = bold   # 실제 문서처럼 번호 부분만 굵게
        p.add_run(rest)

    def bullet(self, text, ilvl=0):
        _number(self.d.add_paragraph(text, style="List Paragraph"), ilvl, num_id=7)

    def save(self, tmp_path):
        path = tmp_path / "cn.docx"
        self.d.save(path)
        return path


@pytest.fixture
def sample(tmp_path):
    d = Doc()
    d.heading("Opening", 0)
    d.heading("Call for IPR", 0, style="Heading 2", num_id=3)  # 번호 레벨이 잘못 매겨진 제목 (무시)
    for title in [ "Agenda", "Highlights", "Minutes", "LS", "Pre-19 LTE", "Pre-19 NR", "Rel-19"]:
        d.heading(title, 0)                                   # 1 .. 8
    d.heading("Release 20 NR", 0)                             # 9
    d.heading("Artificial Intelligence (AI)/Machine Learning (ML)", 1)  # 9.1
    d.heading("CSI compression", 2)                           # 9.1.1
    d.tdoc("R1-2600001", "\tDiscussion on CSI\tVendorA")           # 일반 기고 목록 (무시)
    d.tdoc("R1-2600002", "\tFL summary 1 on 9.1.1\tVendorB", bold=True)
    d.para("Agreement:", bold=True)
    d.para("Support X for inference", bold=True)
    d.bullet("Option A")
    d.bullet("FFS: value of N", ilvl=1)
    d.para("Agreement", bold=True)
    d.para("Support Y")
    d.tdoc("R1-2600003", "\tFL summary 2 on 9.1.1\tVendorB", bold=True)
    d.para("Working assumptions:", bold=True)
    d.para("Z is used")
    d.para("Observation:", bold=True)                         # 합의가 아니므로 제외
    d.para("Something observed")
    # 실제 문서처럼 구간마다 다른 번호 목록을 써도 Agenda 번호는 이어집니다.
    d.heading("Rel-20 Study of 6GR", 0, num_id=2)             # 10
    d.heading("Tasks", 1, num_id=2)                           # 10.1
    d.heading("Multi-antenna system", 1, num_id=2)            # 10.2
    d.heading("CSI acquisition and report", 2, style="Heading 3", num_id=4)  # 10.2.1
    d.heading("Beam management for downlink and uplink", 3, style="Heading 4", num_id=4)  # 10.2.1.1
    d.tdoc("R1-2600005", "\tFL summary on 10.9.9 BM\tVendorC", bold=True)  # 번호 불일치 → 경고
    d.tdoc("R1-2600004", "\tFL summary on 10.2.1.1 BM\tVendorC", bold=True)
    d.para("Conclusion:", bold=True)
    d.para("No consensus on W")
    d.para("Agreement (RAN1#124bis)", bold=True)              # 이전 회의 합의 인용 (신규 아님)
    d.para("Old text")
    d.heading("Unrelated topic", 0)                           # 11
    d.para("Agreement:", bold=True)
    d.para("Not tracked")
    return d.save(tmp_path)


def parse(path):
    return pa.parse_chair_notes(path, "RAN1#999", URL)[0]


def test_warns_when_fl_summary_agenda_number_disagrees(sample):
    _, warnings = pa.parse_chair_notes(sample, "RAN1#999", URL)
    assert len(warnings) == 1
    assert "R1-2600005" in warnings[0] and "10.9.9" in warnings[0] and "10.2.1.1" in warnings[0]


def test_agenda_numbers_follow_word_numbering_levels(sample):
    entries = parse(sample)
    assert [e["agenda_item"] for e in entries[:3]] == ["9.1.1", "9.1.1", "9.1.1"]
    assert entries[3]["agenda_item"] == "10.2.1.1"


def test_extracts_only_decisions_with_types(sample):
    types = [e["type"] for e in parse(sample)]
    assert types == ["Agreement", "Agreement", "Working Assumption", "Conclusion", "Agreement"]


def test_block_text_keeps_bullets_and_stops_at_next_marker(sample):
    first = parse(sample)[0]
    assert first["text"] == "Support X for inference\n- Option A\n  - FFS: value of N"
    assert first["ffs"] == ["FFS: value of N"]
    assert "Support Y" not in first["text"]


def test_source_cites_chair_notes_and_fl_summary(sample):
    entries = parse(sample)
    assert entries[0]["source"]["url"] == URL
    assert entries[0]["source"]["kind"] == "chairman_note"
    assert entries[0]["source"]["tdoc"] == "R1-2600002"
    assert entries[0]["source"]["tdoc_url"].endswith("/TSGR1_999/Docs/R1-2600002.zip")
    assert entries[2]["source"]["tdoc"] == "R1-2600003"


def test_agenda_path_records_every_level_with_number_and_title(sample):
    bm = parse(sample)[3]
    assert bm["agenda_path"] == [
        {"number": "10", "title": "Rel-20 Study of 6GR"},
        {"number": "10.2", "title": "Multi-antenna system"},
        {"number": "10.2.1", "title": "CSI acquisition and report"},
        {"number": "10.2.1.1", "title": "Beam management for downlink and uplink"},
    ]
    assert bm["agenda_title"] == "Beam management for downlink and uplink"
    assert "item" not in bm


def test_decisions_under_every_agenda_are_kept_but_quotes_and_observations_are_not(sample):
    entries = parse(sample)
    texts = " ".join(e["text"] for e in entries)
    assert "Not tracked" in texts                       # 관심 분야 구분 없이 모든 Agenda의 결정을 기록
    assert entries[-1]["agenda_item"] == "11"
    assert "Old text" not in texts
    assert "Something observed" not in texts


def test_ids_are_stable_across_runs(sample):
    assert [e["id"] for e in parse(sample)] == [e["id"] for e in parse(sample)]
    assert parse(sample)[0]["id"] == "RAN1#999:9.1.1:001"


def test_merge_replaces_only_that_meeting():
    db = {"schema_version": 1, "meetings": {}, "entries": [
        {"id": "RAN1#998:1:001", "meeting": "RAN1#998"},
        {"id": "RAN1#999:1:001", "meeting": "RAN1#999"},
    ]}
    new = [{"id": "RAN1#999:2:001", "meeting": "RAN1#999"}]
    pa.merge_meeting(db, "RAN1#999", new, {"ftp_dir": "TSGR1_999"})
    assert [e["id"] for e in db["entries"]] == ["RAN1#998:1:001", "RAN1#999:2:001"]
    assert db["meetings"]["RAN1#999"]["ftp_dir"] == "TSGR1_999"


@pytest.mark.parametrize("folder,meeting", [
    ("TSGR1_126", "RAN1#126"), ("TSGR1_126b", "RAN1#126bis"), ("TSGR1_110b-e", "RAN1#110bis-e"),
])
def test_meeting_id_from_ftp_folder(folder, meeting):
    assert pa.meeting_id_from_folder(folder) == meeting


def test_heading_numbered_shallower_than_its_style_is_still_agenda(tmp_path):
    # 실제 문서의 AIoT 구간: Heading 4에 번호 레벨 2 → 9.1.1 이어야 함
    d = Doc()
    for _ in range(9):
        d.heading("x", 0)
    d.heading("Artificial Intelligence (AI)/Machine Learning (ML)", 1)
    d.heading("Sub topic", 2, style="Heading 4")
    d.para("[126-R20-AI] Email discussion on AI – Moderator")
    d.para("Agreement")
    d.para("Text")
    (e,) = parse(d.save(tmp_path))
    assert e["agenda_item"] == "9.1.1"
    assert e["source"]["discussion"] == "126-R20-AI"
    assert "tdoc" not in e["source"]


def test_official_agenda_numbers_override_word_numbering(tmp_path):
    # 실제 RAN1#124: 공식 Agenda에 '10.5.0 General aspects'가 있어 Word 번호와 하나씩 어긋남
    d = Doc()
    d.heading("Opening of the meeting (Day 1: 9:00 am)", 0)
    d.heading("Rel-20 Study of 6GR", 0)
    d.heading("Multi-antenna system", 1)
    d.heading("General aspects and frameworks", 2, style="Heading 3")
    d.heading("MIMO operation", 2, style="Heading 3")
    d.heading("Beam  management for downlink and uplink ", 3, style="Heading 4")
    d.tdoc("R1-2600009", "\tFL summary on 10.5.1.4\tVendorC", bold=True)
    d.para("Agreement")
    d.para("Text")
    agenda = {
        "1": "Opening of the meeting",
        "10": "Rel-20 Study of 6GR",
        "10.5": "Multi-antenna system",
        "10.5.0": "General aspects and frameworks",
        "10.5.1": "MIMO operation",
        "10.5.1.4": "Beam management for downlink and uplink",
    }
    entries, warnings = pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL, agenda=agenda)
    assert entries[0]["agenda_item"] == "10.5.1.4"
    assert entries[0]["id"] == "RAN1#999:10.5.1.4:001"
    assert warnings == []


def test_heading_missing_from_official_agenda_falls_back_with_warning(tmp_path):
    d = Doc()
    d.heading("Rel-20 Study of 6GR", 0)
    d.heading("Unlisted topic", 1)
    d.para("Agreement")
    d.para("Text")
    entries, warnings = pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL,
                                             agenda={"10": "Rel-20 Study of 6GR"})
    assert entries[0]["agenda_item"] == "10.1"
    assert any("Unlisted topic" in w for w in warnings)


def test_parse_agenda_csv():
    text = '﻿"1","Opening of the meeting"\n"10.5.0","General aspects and frameworks"\n'
    assert pa.parse_agenda_csv(text) == {"1": "Opening of the meeting", "10.5.0": "General aspects and frameworks"}


def test_title_wording_difference_is_fine_when_number_exists_in_agenda(tmp_path):
    # 실제 RAN1#124bis: Chair notes 'ISAC' / 공식 Agenda 'Integrated Sensing And Communication (ISAC)'
    d = Doc()
    d.heading("Rel-20 Study of 6GR", 0)
    d.heading("ISAC", 1)
    d.para("Agreement")
    d.para("Text")
    agenda = {"1": "Rel-20 Study of 6GR", "1.1": "Integrated Sensing And Communication (ISAC)"}
    entries, warnings = pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL, agenda=agenda)
    assert entries[0]["agenda_item"] == "1.1"
    assert warnings == []


def test_marker_variants_seen_in_real_documents_are_read_as_decisions(tmp_path):
    # 실제 문서에서 본 형태: 같은 줄 본문, 세미콜론, 콜론 중복, 설명 붙은 표시, 괄호 설명, 문단 안 줄바꿈
    d = Doc()
    d.heading("Rel-20 Study of 6GR", 0)
    d.heading("Waveform", 1)
    d.para("Agreement: Introduce the following Rel. 19 UE FGs")
    d.para("FFS: component 3")
    d.para("Agreement;")
    d.para("Text after semicolon marker")
    d.para("Agreement:: Adopt the following changes")
    d.para("Conclusion for UL SPS with multi-TBS (after offline):")
    d.para("No consensus")
    d.para("Agreement (updated on Wednesday)")
    d.para("Adopt the TP")
    d.para("Agreement:\nAt least the following directions are studied")
    d.para("Working Assumption from RAN1#124 is not necessarily to be confirmed.")  # 결정 표시가 아님 (본문)
    d.para("Agreements from RAN1#124 on waveform:")                                # 이전 회의 합의 인용
    d.para("Old text")
    stats = {}
    entries, _ = pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL, stats=stats)
    assert [(e["type"], e["text"]) for e in entries] == [
        ("Agreement", "Introduce the following Rel. 19 UE FGs\nFFS: component 3"),
        ("Agreement", "Text after semicolon marker"),
        ("Agreement", "Adopt the following changes"),
        ("Conclusion", "No consensus"),
        ("Agreement", "Adopt the TP"),
        ("Agreement", "At least the following directions are studied\n"
                      "Working Assumption from RAN1#124 is not necessarily to be confirmed."),
    ]
    assert entries[0]["ffs"] == ["FFS: component 3"]
    assert stats["markers"] == 7 and stats["quoted"] == 1 and stats["empty"] == 0
    assert stats["unrecognized"] == []


def test_marker_inside_a_decision_sub_list_stays_in_that_decision(tmp_path):
    d = Doc()
    d.heading("Rel-20 Study of 6GR", 0)
    d.heading("Channel coding", 1)
    d.para("Agreement:")
    d.bullet("For 6G LDPC")
    d.bullet("Working assumption: For data rate within NR range, reuse NR LDPC", ilvl=1)
    d.bullet("FFS: extension", ilvl=1)
    stats = {}
    entries, _ = pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL, stats=stats)
    assert len(entries) == 1 and entries[0]["type"] == "Agreement"
    assert "  - Working assumption: For data rate within NR range, reuse NR LDPC" in entries[0]["text"]
    assert stats["unrecognized"] == []


def test_plain_tdoc_line_right_after_marker_is_the_decision_body(tmp_path):
    d = Doc()
    d.heading("Maintenance", 0)
    d.heading("UE features", 1)
    d.para("Agreement")
    d.tdoc("R1-2509494", " is endorsed for updating RAN1 UE features list.")
    d.para("Agreement")
    d.tdoc("R1-2603365", "\tUpdated RAN1 UE features list\tModerators")
    d.tdoc("R1-2603366", "\tDraft LS on updated features list\tModerators")
    d.tdoc("R1-2603255", "\tSummary#1 of UE features\tModerator", bold=True)   # 다음 결정의 근거 문서
    d.para("Agreement:")
    d.para("Adopt X")
    d.tdoc("R1-2600009", "\tDiscussion on Y\tVendorA")                           # 본문 뒤 TDoc 줄은 결정을 닫음
    d.para("Agreement:")                                                          # 본문 없음
    stats = {}
    entries, _ = pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL, stats=stats)
    assert [e["text"] for e in entries] == [
        "R1-2509494 is endorsed for updating RAN1 UE features list.",
        "R1-2603365\tUpdated RAN1 UE features list\tModerators\nR1-2603366\tDraft LS on updated features list\tModerators",
        "Adopt X",
    ]
    assert entries[2]["source"]["tdoc"] == "R1-2603255"
    assert stats["empty"] == 1


def test_stats_flag_marker_like_lines_the_parser_still_cannot_read(tmp_path):
    d = Doc()
    d.heading("Rel-20 Study of 6GR", 0)
    d.heading("Waveform", 1)
    d.para("Agreement:")
    d.para("Normal decision")
    d.para("Conclusion for case A; case B:")       # 표시처럼 보이지만 형식이 다름
    stats = {}
    pa.parse_chair_notes(d.save(tmp_path), "RAN1#999", URL, stats=stats)
    assert [(u["agenda"], u["text"]) for u in stats["unrecognized"]] == [("1.1", "Conclusion for case A; case B:")]


def test_stats_count_quoted_old_agreements(sample):
    stats = {}
    pa.parse_chair_notes(sample, "RAN1#999", URL, stats=stats)
    assert stats["quoted"] == 1 and stats["unrecognized"] == []
